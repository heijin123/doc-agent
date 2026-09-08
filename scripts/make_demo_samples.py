"""生成 docx/xlsx/pptx/pdf 演示样本（doc-agent 切片验证用）——运行一次即可。

产出去向：data/samples/generated/
- docx_spec.docx          Word 样本：标题层级 + 表格 + 一张内联示意图
- xlsx_products.xlsx      Excel 样本：商品表含空行（验证一行一记录 + 空行跳过）
- pptx_intro.pptx         PPT 样本：两页商务页，第二页含一张图片
- pdf_with_image.pdf      PDF 样本：一页文字 + 一张嵌入图（验证版面图提取）

示意图片用 pymupdf 程序化生成（纯色底 + 英文文字，避开 PDF 中文字体依赖），
无需额外图像依赖；图被 docx/pptx/pdf 三处复用同一字节 → image_id 相同，
可验证"同图跨文档去重由 doc_id 前缀区分"。

依赖：python-docx / openpyxl（pandas 直接读）/ python-pptx / pymupdf——已在 pyproject 声明。
"""
from __future__ import annotations

import io
from pathlib import Path

import pymupdf
from docx import Document
from docx.shared import Inches
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches as PpInches

OUTPUT_DIR = Path(__file__).resolve().parents[1] / "data" / "samples" / "generated"

_PNG_BYTES: bytes | None = None


def _sample_png_bytes() -> bytes:
    """生成一张 320x180 示意 PNG（缓存模块级，docx/pptx/pdf 复用同一字节）。"""
    global _PNG_BYTES
    if _PNG_BYTES is None:
        sample_document = pymupdf.open()
        page = sample_document.new_page(width=320, height=180)
        page.draw_rect(pymupdf.Rect(0, 0, 320, 180), color=(0.2, 0.4, 0.8), fill=(0.92, 0.95, 1.0))
        page.draw_rect(pymupdf.Rect(20, 20, 300, 160), color=(0.2, 0.4, 0.8), width=2)
        page.insert_text((55, 88), "DocAgent", fontsize=26, color=(0.1, 0.2, 0.5))
        page.insert_text((45, 132), "sample image", fontsize=14, color=(0.2, 0.3, 0.6))
        _PNG_BYTES = page.get_pixmap().tobytes("png")
        sample_document.close()
    return _PNG_BYTES


def make_docx_sample() -> None:
    document = Document()
    document.add_heading("产品规格说明", level=1)
    document.add_paragraph("本文件描述 DocAgent 演示产品的完整规格。")
    document.add_heading("核心参数", level=2)
    document.add_paragraph("所有参数均在出厂前经过严格测试。")

    table = document.add_table(rows=3, cols=3)
    table.rows[0].cells[0].text = "型号"
    table.rows[0].cells[1].text = "价格"
    table.rows[0].cells[2].text = "备注"
    table.rows[1].cells[0].text = "DA-100"
    table.rows[1].cells[1].text = "199 元"
    table.rows[1].cells[2].text = "入门款"
    table.rows[2].cells[0].text = "DA-500"
    table.rows[2].cells[1].text = "399 元"
    table.rows[2].cells[2].text = "旗舰款"
    document.add_paragraph("DA-500 支持本地知识库离线问答。")

    document.add_heading("产品示意图", level=2)
    document.add_picture(io.BytesIO(_sample_png_bytes()), width=Inches(3))
    document.add_paragraph("图为 DA-500 外观示意，用于验证 docx 内联图提取。")
    document.save(OUTPUT_DIR / "docx_spec.docx")


def make_xlsx_sample() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "商品清单"
    sheet.append(["商品名称", "价格", "库存"])
    sheet.append(["蓝牙耳机", 199, 50])
    sheet.append(["机械键盘", 399, 30])
    sheet.append([])  # 全空行：验证切块跳过
    sheet.append(["USB-C 扩展坞", 129, 80])
    workbook.save(OUTPUT_DIR / "xlsx_products.xlsx")


def make_pptx_sample() -> None:
    presentation = Presentation()
    first_slide_layout = presentation.slide_layouts[1]  # 标题 + 正文版式

    first_slide = presentation.slides.add_slide(first_slide_layout)
    first_slide.shapes.title.text = "DocAgent 介绍"
    first_slide.placeholders[1].text = "把企业散落的文档变成可检索、可问答、带出处的知识库"

    second_slide = presentation.slides.add_slide(first_slide_layout)
    second_slide.shapes.title.text = "核心能力"
    second_slide.placeholders[1].text = "多格式解析\n结构感知切片\n引用溯源问答"
    second_slide.shapes.add_picture(io.BytesIO(_sample_png_bytes()), PpInches(1), PpInches(2.5), width=PpInches(4))
    presentation.save(OUTPUT_DIR / "pptx_intro.pptx")


def make_pdf_sample() -> None:
    """造一页含嵌入图的 PDF：文字 + 图，验证版面块排序插图提取。"""
    document = pymupdf.open()
    page = document.new_page(width=595, height=842)  # A4 纵向
    page.insert_text((72, 72), "DocAgent Sample PDF with Image", fontsize=16)
    page.insert_text((72, 100), "This page embeds one image below to verify image extraction from layout blocks.", fontsize=11)
    page.insert_text((72, 460), "Image caption: a sample diagram produced by pymupdf.", fontsize=10)
    page.insert_image(pymupdf.Rect(72, 130, 520, 430), stream=_sample_png_bytes())
    document.save(OUTPUT_DIR / "pdf_with_image.pdf")
    document.close()


if __name__ == "__main__":
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    make_docx_sample()
    make_xlsx_sample()
    make_pptx_sample()
    make_pdf_sample()
    print("样本已生成到", OUTPUT_DIR)
